"""Build Additional file 2 (Supplementary Tables S1-S17) for the JHP revision from results/jhp_revision."""
import json
import sys
from importlib.metadata import version, PackageNotFoundError
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "results" / "jhp_revision"
OUT_DIR = ROOT / "Journal of Headache and Pain" / "Revision"
OUT = OUT_DIR / "Additional file 2 - Supplementary Tables.xlsx"


def r(name, **kw):
    return pd.read_csv(RES / name, **kw)


RELABEL = [
    ("(submitted pipeline otherwise)", "(original pipeline otherwise)"),
    ("(submitted pipeline)", "(original pipeline)"),
    ("n_genes_as_submitted", "n_genes_as_scored"),
    ("submitted_abs_cohens_d", "initial_abs_cohens_d"),
    ("submitted_cohens_d", "initial_cohens_d"),
    ("ARI_vs_submitted", "ARI_vs_kmeans_partition"),
    ("submitted_exact_", "initial_exact_"),
    ("legacy_", "initial_"),
    ("revised_", "corrected_"),
]


def relabel(v):
    """Neutral names for the initial (mislabelled-contrast) and corrected module variants."""
    if not isinstance(v, str):
        return v
    for old, new in RELABEL:
        v = v.replace(old, new)
    return v


def stack(blocks):
    """Stack several tables vertically in one sheet, each preceded by a title row."""
    parts = []
    for title, df in blocks:
        head = pd.DataFrame([[title] + [""] * (df.shape[1] - 1)], columns=df.columns)
        cols = pd.DataFrame([list(df.columns)], columns=df.columns)
        blank = pd.DataFrame([[""] * df.shape[1]], columns=df.columns)
        parts += [head, cols, df.astype(object), blank]
    width = max(p.shape[1] for p in parts)
    out = []
    for p in parts:
        p = p.copy()
        p.columns = range(p.shape[1])
        out.append(p.reindex(columns=range(width)))
    return pd.concat(out, ignore_index=True)


def sample_table():
    tech = r("B7_technical_covariates_all20.csv", index_col=0).drop(columns=["mito_fraction"])
    sex = r("B7_sex_check.csv", index_col=0)[["sex_GEO", "Ychr_mean", "sex_predicted", "concordant"]]
    proj = r("B5_controls_projected_onto_TN_partition.csv").set_index("sample")["assigned_subtype"]
    pcs = r("B7_TN_PC_scores.csv", index_col=0)[["PC1", "PC2"]]
    t = tech.join(sex).join(pcs)
    t["partition_or_projection"] = t["group"]
    for s, a in proj.items():
        t.loc[s, "partition_or_projection"] = f"Control (projected to Subtype {a})"
    t.index.name = "sample"
    t = t.rename(columns={
        "group": "group", "genes_detected_FPKM_gt_0.1": "genes_detected_FPKM>0.1",
        "top100_gene_fraction": "fraction_of_FPKM_in_top100_genes",
        "globin_fraction_of_total_FPKM": "globin_fraction_of_total_FPKM",
        "Erythroid": "erythroid_marker_score", "Platelet": "platelet_marker_score",
        "Neutrophil": "neutrophil_marker_score", "Ychr_mean": "mean_log2FPKM_Y_genes",
        "sex_predicted": "sex_from_expression", "concordant": "sex_concordant",
        "PC1": "TN_PC1", "PC2": "TN_PC2"})
    order = ["group", "partition_or_projection", "sex_GEO", "sex_from_expression", "sex_concordant",
             "mean_log2FPKM_Y_genes", "genes_detected_FPKM>0.1", "total_FPKM",
             "fraction_of_FPKM_in_top100_genes", "globin_fraction_of_total_FPKM",
             "erythroid_marker_score", "platelet_marker_score", "neutrophil_marker_score", "TN_PC1", "TN_PC2"]
    return t[order].sort_values(["group", "sample"]).reset_index()


def group_summary():
    t = sample_table()
    rows = []
    for g, d in [("All TN (n=10)", t[t.group.str.startswith("Subtype")]),
                 ("Subtype 0 (n=6)", t[t.group == "Subtype 0"]),
                 ("Subtype 1 (n=4)", t[t.group == "Subtype 1"]),
                 ("Controls (n=10)", t[t.group == "Control"])]:
        rows.append({"group": g, "female_n": int((d.sex_GEO == "F").sum()), "male_n": int((d.sex_GEO == "M").sum()),
                     "age": "not deposited in GEO", "medication": "not deposited in GEO",
                     "median_globin_fraction": round(d.globin_fraction_of_total_FPKM.median(), 3),
                     "median_genes_detected": int(d["genes_detected_FPKM>0.1"].median())})
    return pd.DataFrame(rows)


def software_versions():
    pk = ["numpy", "pandas", "scipy", "scikit-learn", "statsmodels", "pydeseq2", "gseapy", "matplotlib",
          "python-docx", "openpyxl", "requests"]
    rows = [{"software": "Python", "version": sys.version.split()[0], "use": "all analyses"}]
    use = {"numpy": "numerics", "pandas": "data handling", "scipy": "statistics, NNLS, lowess inputs",
           "scikit-learn": "PCA, k-means, hierarchical clustering, silhouette, Calinski-Harabasz, ARI",
           "statsmodels": "lowess trend for limma-trend prior, multiple testing",
           "pydeseq2": "IoN-CCI count-based differential expression (and GSE186505 salmon counts)",
           "gseapy": "Enrichr libraries, preranked GSEA", "matplotlib": "figures",
           "python-docx": "manuscript assembly", "openpyxl": "tables", "requests": "data retrieval"}
    for p in pk:
        try:
            rows.append({"software": p, "version": version(p), "use": use.get(p, "")})
        except PackageNotFoundError:
            pass
    rows += [
        {"software": "salmon", "version": "1.10.3", "use": "re-quantification of GSE186505 raw reads (GENCODE v26 index)"},
        {"software": "limma-trend (re-implemented)", "version": "Smyth 2004; Law et al. 2014",
         "use": "moderated t-statistics on log2(FPKM+1); implementation in scripts/run_jhp_revision_analyses.py"},
        {"software": "ABIS RNA-seq signature", "version": "Monaco et al. 2019",
         "use": "reference-based NNLS deconvolution (relative estimates only)"},
        {"software": "Enrichr libraries", "version": "Reactome_2022; GO_Biological_Process_2023", "use": "ORA and GSEA"},
        {"software": "MGI HOM_MouseHumanSequence.rpt", "version": "downloaded 2026-09", "use": "one-to-one mouse-human orthologs"},
    ]
    return pd.DataFrame(rows)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    comp = r("B6_cell_composition_tests.csv")

    sheets = {}
    sheets["S1 Participants"] = stack([
        ("Table S1a. Group-level characteristics available from the GEO deposit (GSE186505).", group_summary()),
        ("Table S1b. Per-sample characteristics, expression-based sex check and library-composition covariates.", sample_table())])
    sheets["S2 IoN-CCI label audit"] = r("A1_ion_cci_sample_label_audit.csv")
    sheets["S3 IoN-CCI DE TG"] = r("A2_ion_cci_deseq2_tg_injury_vs_sham.csv").rename(columns={"Unnamed: 0": "ensembl_id"})
    sheets["S4 IoN-CCI DE Sp5C"] = r("A2_ion_cci_deseq2_sp5c_injury_vs_sham.csv").rename(columns={"Unnamed: 0": "ensembl_id"})
    sheets["S5 IoN-CCI tissue identity"] = r("A2_ion_cci_deseq2_sham_TG_vs_Sp5C_tissue_identity.csv").rename(columns={"Unnamed: 0": "ensembl_id"})
    sheets["S6 IoN-CCI Reactome ORA"] = r("A3_ion_cci_reactome_ORA_by_tissue_direction.csv").sort_values(["tissue", "direction", "p_adj"])
    defs = json.loads((RES / "A4_revised_module_definitions.json").read_text())
    defs_df = pd.DataFrame([{"module": k, **({"definition": v} if isinstance(v, str) else v)} for k, v in defs.items()])
    sheets["S7 Module gene sets"] = stack([
        ("Table S7a. Ortholog coverage and whole-blood expression of initial and corrected modules.",
         r("B4_module_ortholog_blood_expression.csv")),
        ("Table S7b. Corrected module definitions.", defs_df.astype(str)),
        ("Table S7c. Corrected module gene lists (human one-to-one orthologs).", r("A4_revised_module_gene_sets.csv"))])
    sheets["S8 Clusterability"] = stack([
        ("Table S8a. k-selection metrics for k=2-4 with single-Gaussian null reference (1,000 null datasets).",
         r("B2_k_selection_metrics_with_null.csv")),
        ("Table S8b. Gap statistic, z-scored FPKM (Tibshirani 1-SE rule).", r("B2_gap_statistic_k1_4.csv")),
        ("Table S8c. Gap statistic, z-scored log2(FPKM+1).", r("B2_gap_statistic_k1_4_log2.csv"))])
    sheets["S9 Stability"] = stack([
        ("Table S9a. Leave-one-out re-clustering.", r("B3_leave_one_out.csv")),
        ("Table S9b. Cluster-wise Jaccard stability (subsampling 8 of 10 and bootstrap).", r("B3_clusterwise_jaccard.csv")),
        ("Table S9c. Analytical sensitivity variants.", r("B3_sensitivity_variants.csv")),
        ("Table S9d. Consensus matrix (subsampling).", r("B3_consensus_matrix_subsampling.csv").rename(columns={"Unnamed: 0": "sample"}))])
    sheets["S10 Circularity nulls"] = stack([
        ("Table S10a. Reproduction of initial module effect sizes.", r("B4_legacy_effect_reproduction.csv")),
        ("Table S10b. Initial gene sets (as originally scored) under the full-pipeline null.", r("B4b_submitted_gene_sets_full_pipeline_null.csv")),
        ("Table S10c. Initial gene sets (as originally scored) against 10,000 expression-matched random gene sets.",
         r("B4c_submitted_gene_sets_matched_random_null.csv")),
        ("Table S10d. Blood-expressed module genes: exact permutation, full-pipeline null and matched random gene-set null.",
         r("B4_module_effects_circularity_nulls.csv"))]
        + ([("Table S10e. Sp5C-down ECM/Schwann module (added after label correction; analysed separately from the "
             "pre-specified gene sets): definition, blood expression, partition effect with full-pipeline and random "
             "gene-set nulls, globin depletion, and TN versus control mean and dispersion.",
             r("E_sp5c_ecm/E3_sp5c_ecm_tests.csv").rename(columns={"Unnamed: 0": "statistic"})),
            ("Table S10f. Sp5C-down ECM/Schwann module genes and their mean FPKM in blood.",
             r("E_sp5c_ecm/E1_sp5c_ecm_module_genes.csv")),
            ("Table S10g. Sp5C-down ECM/Schwann module scores per sample (injury-oriented; residual after regression "
             "on globin fraction).", r("E_sp5c_ecm/E2_sp5c_ecm_scores_all20.csv"))]
           if (RES / "E_sp5c_ecm" / "E3_sp5c_ecm_tests.csv").exists() else []))
    sheets["S11 Controls"] = stack([
        ("Table S11a. Controls projected onto the TN partition.", r("B5_controls_projected_onto_TN_partition.csv")),
        ("Table S11b. Module scores: controls versus partitions.", r("B5_module_scores_controls_vs_subtypes.csv")),
        ("Table S11c. Per-sample module scores (mean z of log2(FPKM+1), injury-oriented).",
         r("B5_module_scores_all20_log2z.csv").rename(columns={"Unnamed: 0": "sample"}))]
        + ([("Table S11d. Dispersion, TN versus controls (diagnosis labels only): globin fraction, first two PCs of all "
             "20 samples and 11 distinct module scores. Brown–Forsythe test and permutation test of the variance ratio "
             "(20,000 permutations); Benjamini–Hochberg q across the 14 features.", r("D_technical/D4_dispersion_TN_vs_control.csv")),
            ("Table S11e. Module-score dispersion after regressing each score on globin fraction (all 20 samples); "
             "Spearman correlation and R² of each score with globin fraction.", r("D_technical/D5_module_dispersion_globin_adjusted.csv"))]
           if (RES / "D_technical" / "D5_module_dispersion_globin_adjusted.csv").exists() else []))
    sheets["S12 Cell composition"] = stack([
        ("Table S12a. Composition tests (ABIS NNLS fractions and marker scores).", comp),
        ("Table S12b. ABIS NNLS estimates per sample (relative, not absolute).", r("B6_ABIS_NNLS_deconvolution_all20.csv")),
        ("Table S12c. Marker-gene scores per sample.", r("B6_marker_scores_all20.csv").rename(columns={"Unnamed: 0": "sample"}))])
    sheets["S13 Composition confounding"] = stack([
        ("Table S13a. Technical and composition covariates versus partition and principal components (TN).",
         r("B7_confounder_tests_TN.csv").dropna(subset=["subtype_MWU_exact_p"])),
        ("Table S13b. Globin fraction per sample.", r("B10_globin_fraction_by_sample.csv").rename(columns={"Unnamed: 0": "sample"})),
        ("Table S13c. Partition after globin depletion or residualisation.", r("B10_partition_after_globin_adjustment.csv")),
        ("Table S13d. Module effect sizes before and after globin depletion.", r("B10_module_effects_after_globin_depletion.csv")),
        ("Table S13e. Top 200 PC1 loadings (TN).", r("B7_PC1_top200_loadings.csv"))]
        + ([("Table S13f. Per-sample technical covariates from the raw reads: read pairs, salmon mapping rate, "
             "sequencing run (instrument:run:flowcell:lane from FASTQ headers) and flowcell batch.",
             r("D_technical/D1_technical_covariates_per_sample.csv")),
            ("Table S13g. Technical covariates versus partition, diagnosis, PC1 and globin fraction.",
             r("D_technical/D2_technical_covariates_tests.csv"))]
           if (RES / "D_technical" / "D2_technical_covariates_tests.csv").exists() else []))
    sheets["S14 DE TN vs control"] = stack([
        ("Table S14a. TN versus control, limma-trend moderated t, sex-adjusted (11,219 expressed genes).",
         r("B8_TN_vs_control_limma_trend_sex_adjusted.csv")),
        ("Table S14b. Preranked GSEA on the moderated t (Reactome 2022, GO BP 2023), FDR < 0.25.",
         r("B9_prerank_GSEA_TN_vs_control_moderated_t.csv").query("`FDR q-val` < 0.25").sort_values("FDR q-val"))])
    sheets["S15 DE partition (descr.)"] = stack([
        ("Table S15a. Subtype 1 versus Subtype 0, limma-trend, sex-adjusted. Descriptive only: the partition was derived "
         "from the same data, so these P values are not valid inference.", r("B8_subtype1_vs_0_limma_trend_sex_adjusted.csv")),
        ("Table S15b. Preranked GSEA for the partition contrast, FDR < 0.25 (descriptive).",
         r("B9_prerank_GSEA_subtype1_vs_0_moderated_t.csv").query("`FDR q-val` < 0.25").sort_values("FDR q-val"))])
    salmon = RES / "C_salmon"
    if (salmon / "C_summary_tables.csv").exists():
        blocks = [("Table S16a. Re-quantification from raw reads (salmon): summary.", pd.read_csv(salmon / "C_summary_tables.csv"))]
        for f, title in [("C1_salmon_qc.csv", "Table S16b. Per-sample mapping rate, agreement with deposited FPKM and globin fraction."),
                         ("C2_TN_vs_control_DESeq2_sex_adjusted.csv", "Table S16c. TN versus control, PyDESeq2 on salmon counts (~ sex + condition)."),
                         ("C2_TN_vs_control_DESeq2_sex_globin_adjusted.csv", "Table S16d. TN versus control, PyDESeq2 on salmon counts (~ sex + globin fraction + condition)."),
                         ("C2_TN_vs_control_DESeq2_sex_adjusted_globin_genes_removed.csv", "Table S16e. TN versus control, PyDESeq2 on salmon counts with haemoglobin genes removed (~ sex + condition)."),
                         ("C4_TN_vs_control_limma_trend_sex_globin_adjusted.csv", "Table S16f. TN versus control, moderated t on deposited FPKM (~ sex + globin fraction + condition)."),
                         ("C3_partition_on_salmon_counts.csv", "Table S16g. Clustering and globin fraction on salmon-derived expression.")]:
            if (salmon / f).exists():
                blocks.append((title, pd.read_csv(salmon / f)))
        tech = RES / "D_technical"
        for f, title in [("D3_TN_vs_control_limma_trend_sex_batch.csv", "Table S16h. TN versus control, moderated t on deposited FPKM (~ sex + sequencing flowcell + condition)."),
                         ("D3_TN_vs_control_limma_trend_sex_batch_globin.csv", "Table S16i. TN versus control, moderated t on deposited FPKM (~ sex + sequencing flowcell + globin fraction + condition)."),
                         ("D3_TN_vs_control_DESeq2_sex_batch.csv", "Table S16j. TN versus control, PyDESeq2 on salmon counts (~ sex + sequencing flowcell + condition).")]:
            if (tech / f).exists():
                blocks.append((title, pd.read_csv(tech / f)))
        sheets["S16 Raw-read requant"] = stack(blocks)
    sheets["S17 Software"] = software_versions()

    readme = pd.DataFrame([
        ("S1 Participants", "Participant characteristics available in GEO, expression-based sex check, library-composition covariates."),
        ("S2 IoN-CCI label audit", "PRJNA991739 run-to-condition/tissue mapping from SRA metadata versus the labels used in an initial version of the analysis."),
        ("S3 IoN-CCI DE TG", "PyDESeq2, IoN-CCI versus sham within trigeminal ganglion (3 vs 3 pooled libraries)."),
        ("S4 IoN-CCI DE Sp5C", "PyDESeq2, IoN-CCI versus sham within spinal trigeminal nucleus caudalis (3 vs 3)."),
        ("S5 IoN-CCI tissue identity", "PyDESeq2, sham TG versus sham Sp5C, used to diagnose the label error."),
        ("S6 IoN-CCI Reactome ORA", "Reactome 2022 over-representation by tissue and direction (padj < 0.05 genes)."),
        ("S7 Module gene sets", "Module definitions, gene lists, ortholog counts and blood expression."),
        ("S8 Clusterability", "Elbow, silhouette, Calinski-Harabasz with null reference; gap statistic including k = 1."),
        ("S9 Stability", "Leave-one-out, Jaccard, consensus, seed and analytical-choice sensitivity."),
        ("S10 Circularity nulls", "Full-pipeline and random gene-set nulls for module effect sizes; Sp5C-down ECM/Schwann module."),
        ("S11 Controls", "Pain-free controls projected onto the partition and scored on the modules; TN versus control dispersion, with and without globin adjustment."),
        ("S12 Cell composition", "ABIS NNLS deconvolution and marker-gene scores with tests."),
        ("S13 Composition confounding", "Globin/erythroid library composition versus the partition; globin depletion; read depth, mapping rate and sequencing flowcell."),
        ("S14 DE TN vs control", "Moderated-t differential expression and preranked GSEA, TN versus control."),
        ("S15 DE partition (descr.)", "Moderated-t contrast between the partitions (descriptive, circular)."),
        ("S16 Raw-read requant", "Salmon re-quantification of GSE186505 raw reads, count-based analyses and flowcell-adjusted TN versus control models."),
        ("S17 Software", "Software and resource versions."),
    ], columns=["sheet", "content"])
    readme = readme[readme.sheet.isin(sheets.keys())]

    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        readme.to_excel(xw, sheet_name="README", index=False)
        for name, df in sheets.items():
            stacked = isinstance(df.columns, pd.RangeIndex)
            df = df.rename(columns=relabel).map(relabel)
            df.to_excel(xw, sheet_name=name[:31], index=False, header=not stacked)

    wb = load_workbook(OUT)
    bold = Font(bold=True, name="Arial", size=9)
    norm = Font(name="Arial", size=9)
    fill = PatternFill("solid", fgColor="E8EEF4")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = norm
        first = [c.value for c in ws[1]]
        if isinstance(first[0], str) and first[0].startswith("Table S"):
            for row in ws.iter_rows():
                v = row[0].value
                if isinstance(v, str) and v.startswith("Table S"):
                    row[0].font = Font(bold=True, name="Arial", size=10)
                    row[0].alignment = Alignment(wrap_text=False)
        else:
            for c in ws[1]:
                c.font = bold
                c.fill = fill
            ws.freeze_panes = "A2"
        for i, col in enumerate(ws.columns, 1):
            vals = [len(str(c.value)) for c in list(col)[:200] if c.value is not None and not str(c.value).startswith("Table S")]
            ws.column_dimensions[get_column_letter(i)].width = min(max(vals + [8]) + 2, 60)
    wb.save(OUT)
    print("saved", OUT, "sheets:", len(wb.sheetnames))


if __name__ == "__main__":
    main()
